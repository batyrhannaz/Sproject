from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db.models import Count, Q
from django.http import HttpResponse
from datetime import timedelta, date
import csv
import io
import openpyxl

from accounts.models import User, Role
from accounts.decorators import role_required
from library.models import (
    LibrarySettings, Genre, Book, BookCopy, Reservation, BorrowRecord, BookRecommendation
)
from school.models import SchoolClass


# ----------------------------------------------------
# 1. CATALOG & BOOK DETAILS & RESERVATION
# ----------------------------------------------------
@login_required
def catalog_view(request):
    """
    Public / reader catalog with search, filters, 'Новинки', and 'Популярное'.
    """
    settings = LibrarySettings.get_settings()
    genres = Genre.objects.all().order_by('name')

    q = request.GET.get('q', '').strip()
    genre_id = request.GET.get('genre')
    grade = request.GET.get('grade')
    lang = request.GET.get('lang')
    available_only = request.GET.get('available') == '1'
    section = request.GET.get('section', 'all')  # all, new, popular

    books = Book.objects.all().select_related('genre').prefetch_related('copies')

    if q:
        books = books.filter(
            Q(title__icontains=q) |
            Q(author__icontains=q) |
            Q(isbn__icontains=q)
        )
    if genre_id:
        books = books.filter(genre_id=genre_id)
    if grade:
        books = books.filter(grade_level=grade)
    if lang:
        books = books.filter(language=lang)
    if available_only:
        books = books.filter(copies__status='available').distinct()

    if section == 'new':
        books = books.filter(Q(is_new=True) | Q(created_at__gte=timezone.now() - timedelta(days=30)))
    elif section == 'popular':
        books = books.annotate(num_borrows=Count('copies__borrow_records')).order_by('-num_borrows')

    # Teacher recommendations
    recommendations = BookRecommendation.objects.select_related('teacher', 'book', 'target_class').order_by('-created_at')[:6]

    return render(request, 'library/catalog.html', {
        'books': books,
        'genres': genres,
        'selected_q': q,
        'selected_genre': genre_id,
        'selected_grade': grade,
        'selected_lang': lang,
        'available_only': available_only,
        'section': section,
        'recommendations': recommendations,
        'grade_choices': range(1, 12),
        'lang_choices': Book.LANG_CHOICES,
    })


@login_required
def book_detail_view(request, book_id):
    """
    Detailed book view with copies list and reservation button.
    """
    book = get_object_or_404(Book.objects.prefetch_related('copies'), id=book_id)
    settings = LibrarySettings.get_settings()
    user = request.user

    # Check if user has overdue books
    today = timezone.now().date()
    has_overdue = BorrowRecord.objects.filter(reader=user, return_date__isnull=True, due_date__lt=today).exists()

    # Active borrows and reservations count for limit
    active_borrows_count = BorrowRecord.objects.filter(reader=user, return_date__isnull=True).count()
    active_reservations_count = Reservation.objects.filter(reader=user, status__in=['pending', 'confirmed']).count()
    total_held = active_borrows_count + active_reservations_count

    already_reserved = Reservation.objects.filter(book=book, reader=user, status__in=['pending', 'confirmed']).exists()

    recommendations = BookRecommendation.objects.filter(book=book).select_related('teacher', 'target_class')

    return render(request, 'library/book_detail.html', {
        'book': book,
        'has_overdue': has_overdue,
        'total_held': total_held,
        'max_limit': settings.max_books_per_reader,
        'can_reserve': not has_overdue and total_held < settings.max_books_per_reader and book.available_copies_count > 0 and not already_reserved,
        'already_reserved': already_reserved,
        'recommendations': recommendations,
    })


@login_required
def reserve_book_view(request, book_id):
    """
    Student or reader creates a reservation for a book.
    Enforces rules:
    - Reader with overdue book cannot reserve
    - Reader cannot exceed max_books_per_reader limit
    - Reservation expires in 3 days (or configured value)
    """
    if request.method != 'POST':
        return redirect('library:book_detail', book_id=book_id)

    book = get_object_or_404(Book, id=book_id)
    settings = LibrarySettings.get_settings()
    user = request.user
    today = timezone.now().date()

    # 1. Check for overdue books
    has_overdue = BorrowRecord.objects.filter(reader=user, return_date__isnull=True, due_date__lt=today).exists()
    if has_overdue:
        messages.error(
            request,
            "Бронирование заблокировано! У вас есть просроченная литература. "
            "Пожалуйста, верните долги в школьную библиотеку."
        )
        return redirect('library:book_detail', book_id=book.id)

    # 2. Check reader quota
    active_borrows = BorrowRecord.objects.filter(reader=user, return_date__isnull=True).count()
    active_res = Reservation.objects.filter(reader=user, status__in=['pending', 'confirmed']).count()
    if (active_borrows + active_res) >= settings.max_books_per_reader:
        messages.error(
            request,
            f"Вы исчерпали лимит книг ({settings.max_books_per_reader} шт.). "
            "Верните ранее взятые книги для нового бронирования."
        )
        return redirect('library:book_detail', book_id=book.id)

    # 3. Check already reserved
    if Reservation.objects.filter(book=book, reader=user, status__in=['pending', 'confirmed']).exists():
        messages.warning(request, "Вы уже забронировали эту книгу.")
        return redirect('library:my_books')

    # 4. Check available copies
    available_copy = book.copies.filter(status='available').first()
    if not available_copy:
        messages.error(request, "К сожалению, в данный момент свободных экземпляров нет.")
        return redirect('library:book_detail', book_id=book.id)

    # Mark copy as reserved
    available_copy.status = 'reserved'
    available_copy.save()

    # Create reservation
    expires_at = timezone.now() + timedelta(days=settings.reservation_validity_days)
    Reservation.objects.create(
        book=book,
        reader=user,
        copy=available_copy,
        status='pending',
        expires_at=expires_at
    )

    messages.success(
        request,
        f"Книга «{book.title}» успешно забронирована! "
        f"Бронь действительна до {expires_at.strftime('%d.%m.%Y %H:%M')}. Обратитесь к библиотекарю для получения."
    )
    return redirect('library:my_books')


@login_required
def cancel_reservation_view(request, reservation_id):
    reservation = get_object_or_404(Reservation, id=reservation_id, reader=request.user)
    if reservation.status in ['pending', 'confirmed']:
        if reservation.copy:
            reservation.copy.status = 'available'
            reservation.copy.save()
        reservation.status = 'cancelled'
        reservation.save()
        messages.info(request, f"Бронь на книгу «{reservation.book.title}» отменена.")
    return redirect('library:my_books')


@login_required
def my_books_view(request):
    """
    Shows reader's active borrowings, past history, and reservations.
    """
    user = request.user
    today = timezone.now().date()

    # If parent, can view child's books as well
    target_user = user
    child_id = request.GET.get('child_id')
    if user.is_parent:
        child = User.objects.filter(id=child_id, parent_relations__parent=user).first()
        if child:
            target_user = child
        else:
            first_rel = user.children_relations.first()
            if first_rel:
                target_user = first_rel.student

    active_borrows = BorrowRecord.objects.filter(
        reader=target_user, return_date__isnull=True
    ).select_related('copy__book').order_by('due_date')

    past_borrows = BorrowRecord.objects.filter(
        reader=target_user, return_date__isnull=False
    ).select_related('copy__book').order_by('-return_date')[:20]

    reservations = Reservation.objects.filter(
        reader=target_user
    ).select_related('book', 'copy').order_by('-created_at')

    children = user.children_relations.all() if user.is_parent else []

    return render(request, 'library/my_books.html', {
        'target_user': target_user,
        'children': children,
        'active_borrows': active_borrows,
        'past_borrows': past_borrows,
        'reservations': reservations,
        'today': today,
    })


# ----------------------------------------------------
# 2. LIBRARIAN CABINET
# ----------------------------------------------------
@login_required
@role_required(['librarian', 'director'])
def librarian_cabinet_view(request):
    """
    Librarian Control Center:
    - Quick issue / checkout
    - Confirm & fulfill reservations
    - Return book & record condition
    - Overdue debtors
    - Catalog and copy management
    - Fund reports & Excel export
    - Mass catalog import
    """
    settings = LibrarySettings.get_settings()
    today = timezone.now().date()

    # Tab selection
    current_tab = request.GET.get('tab', 'borrows')

    # Metrics
    total_books = Book.objects.count()
    total_copies = BookCopy.objects.exclude(status='written_off').count()
    active_borrows_qs = BorrowRecord.objects.filter(return_date__isnull=True).select_related('reader', 'copy__book')
    active_borrows_count = active_borrows_qs.count()
    overdue_borrows_qs = active_borrows_qs.filter(due_date__lt=today).order_by('due_date')
    overdue_count = overdue_borrows_qs.count()
    pending_reservations = Reservation.objects.filter(status='pending').select_related('reader', 'book', 'copy').order_by('-created_at')

    # Active borrows list (with search/filter)
    search_reader = request.GET.get('search_reader', '').strip()
    filtered_borrows = active_borrows_qs
    if search_reader:
        filtered_borrows = filtered_borrows.filter(
            Q(reader__first_name__icontains=search_reader) |
            Q(reader__last_name__icontains=search_reader) |
            Q(copy__book__title__icontains=search_reader) |
            Q(copy__inventory_number__icontains=search_reader)
        )

    # Readers list
    readers = User.objects.filter(role__in=[Role.STUDENT, Role.TEACHER]).order_by('last_name', 'first_name')
    available_copies = BookCopy.objects.filter(status='available').select_related('book').order_by('book__title')
    genres = Genre.objects.all().order_by('name')
    books_catalog = Book.objects.all().prefetch_related('copies').order_by('title')

    # 1. Handle Quick Issue
    if request.method == 'POST' and request.POST.get('action') == 'quick_issue':
        reader_id = request.POST.get('reader_id')
        copy_id = request.POST.get('copy_id')
        loan_days = int(request.POST.get('loan_days') or settings.default_loan_days)

        if reader_id and copy_id:
            reader = get_object_or_404(User, id=reader_id)
            copy = get_object_or_404(BookCopy, id=copy_id)

            if copy.status != 'available':
                messages.error(request, f"Экземпляр {copy.inventory_number} недоступен (статус: {copy.get_status_display()}).")
            else:
                due = today + timedelta(days=loan_days)
                copy.status = 'borrowed'
                copy.save()

                BorrowRecord.objects.create(
                    copy=copy,
                    reader=reader,
                    issued_by=request.user,
                    borrow_date=today,
                    due_date=due,
                    initial_condition=copy.condition
                )
                messages.success(request, f"Книга «{copy.book.title}» (№ {copy.inventory_number}) выдана читателю {reader.display_name} до {due.strftime('%d.%m.%Y')}.")
                return redirect(f"{request.path}?tab=borrows")

    # 2. Handle Confirm / Fulfill Reservation
    elif request.method == 'POST' and request.POST.get('action') == 'fulfill_reservation':
        res_id = request.POST.get('reservation_id')
        reservation = get_object_or_404(Reservation, id=res_id)
        copy = reservation.copy or reservation.book.copies.filter(status__in=['available', 'reserved']).first()

        if copy:
            copy.status = 'borrowed'
            copy.save()
            reservation.status = 'fulfilled'
            reservation.save()

            due = today + timedelta(days=settings.default_loan_days)
            BorrowRecord.objects.create(
                copy=copy,
                reader=reservation.reader,
                issued_by=request.user,
                borrow_date=today,
                due_date=due,
                initial_condition=copy.condition,
                reservation=reservation
            )
            messages.success(request, f"Бронь подтверждена! Книга «{copy.book.title}» выдана {reservation.reader.display_name} до {due.strftime('%d.%m.%Y')}.")
            return redirect(f"{request.path}?tab=reservations")
        else:
            messages.error(request, "Нет доступного экземпляра для выдачи!")

    # 3. Handle Reject Reservation
    elif request.method == 'POST' and request.POST.get('action') == 'reject_reservation':
        res_id = request.POST.get('reservation_id')
        reservation = get_object_or_404(Reservation, id=res_id)
        if reservation.copy:
            reservation.copy.status = 'available'
            reservation.copy.save()
        reservation.status = 'cancelled'
        reservation.save()
        messages.info(request, "Бронь отклонена.")
        return redirect(f"{request.path}?tab=reservations")

    # 4. Handle Return Book
    elif request.method == 'POST' and request.POST.get('action') == 'return_book':
        borrow_id = request.POST.get('borrow_id')
        condition = request.POST.get('return_condition', 'good')
        borrow = get_object_or_404(BorrowRecord, id=borrow_id)

        borrow.return_date = today
        borrow.return_condition = condition
        borrow.save()

        # Update copy status & condition
        copy = borrow.copy
        copy.condition = condition
        copy.status = 'available'
        copy.save()

        messages.success(request, f"Книга «{copy.book.title}» принята от {borrow.reader.display_name}! Состояние: {copy.get_condition_display()}.")
        return redirect(f"{request.path}?tab=borrows")

    # 5. Handle Add New Book
    elif request.method == 'POST' and request.POST.get('action') == 'add_book':
        title = request.POST.get('title')
        author = request.POST.get('author')
        genre_id = request.POST.get('genre')
        grade = request.POST.get('grade_level')
        lang = request.POST.get('language', 'ru')
        year = request.POST.get('publication_year')
        isbn = request.POST.get('isbn', '')
        desc = request.POST.get('description', '')
        cover = request.FILES.get('cover_image')
        copies_to_add = int(request.POST.get('initial_copies') or 1)

        if title and author:
            genre = Genre.objects.filter(id=genre_id).first() if genre_id else None
            book = Book.objects.create(
                title=title, author=author, genre=genre,
                grade_level=int(grade) if grade else None,
                language=lang,
                publication_year=int(year) if year else None,
                isbn=isbn, description=desc, cover_image=cover
            )
            # Create initial copies
            for i in range(1, copies_to_add + 1):
                inv = f"INV-{book.id}-{i:03d}"
                BookCopy.objects.create(book=book, inventory_number=inv, status='available', condition='excellent')

            messages.success(request, f"Книга «{book.title}» успешно добавлена с {copies_to_add} экз.!")
            return redirect(f"{request.path}?tab=catalog")

    # 6. Handle Add Copy to Book
    elif request.method == 'POST' and request.POST.get('action') == 'add_copy':
        book_id = request.POST.get('book_id')
        inv = request.POST.get('inventory_number')
        cond = request.POST.get('condition', 'good')
        book = get_object_or_404(Book, id=book_id)

        if inv:
            if BookCopy.objects.filter(inventory_number=inv).exists():
                messages.error(request, f"Экземпляр с номером {inv} уже существует!")
            else:
                BookCopy.objects.create(book=book, inventory_number=inv, condition=cond, status='available')
                messages.success(request, f"Экземпляр № {inv} добавлен к книге «{book.title}».")
        return redirect(f"{request.path}?tab=catalog")

    # 7. Handle Write-off Copy (Списание)
    elif request.method == 'POST' and request.POST.get('action') == 'write_off_copy':
        copy_id = request.POST.get('copy_id')
        reason = request.POST.get('reason', '')
        copy = get_object_or_404(BookCopy, id=copy_id)
        copy.status = 'written_off'
        copy.notes = f"Списано: {reason}"
        copy.save()
        messages.warning(request, f"Экземпляр № {copy.inventory_number} списан.")
        return redirect(f"{request.path}?tab=catalog")

    # 8. Handle Library Settings Update
    elif request.method == 'POST' and request.POST.get('action') == 'update_settings':
        settings.default_loan_days = int(request.POST.get('default_loan_days', 14))
        settings.max_books_per_reader = int(request.POST.get('max_books_per_reader', 3))
        settings.reservation_validity_days = int(request.POST.get('reservation_validity_days', 3))
        settings.save()
        messages.success(request, "Правила библиотеки успешно обновлены!")
        return redirect(f"{request.path}?tab=rules")

    # Reports Data
    popular_books = Book.objects.annotate(borrow_count=Count('copies__borrow_records')).order_by('-borrow_count')[:10]
    active_readers = User.objects.filter(role__in=[Role.STUDENT, Role.TEACHER]).annotate(
        borrows_total=Count('borrow_records')
    ).order_by('-borrows_total')[:10]

    fund_by_genre = Genre.objects.annotate(copies_count=Count('books__copies')).order_by('-copies_count')
    fund_by_condition = BookCopy.objects.values('condition').annotate(count=Count('id'))

    return render(request, 'library/librarian_cabinet.html', {
        'current_tab': current_tab,
        'settings': settings,
        'today': today,
        'total_books': total_books,
        'total_copies': total_copies,
        'active_borrows_count': active_borrows_count,
        'overdue_count': overdue_count,
        'filtered_borrows': filtered_borrows,
        'overdue_borrows': overdue_borrows_qs,
        'pending_reservations': pending_reservations,
        'readers': readers,
        'available_copies': available_copies,
        'books_catalog': books_catalog,
        'genres': genres,
        'popular_books': popular_books,
        'active_readers': active_readers,
        'fund_by_genre': fund_by_genre,
        'fund_by_condition': fund_by_condition,
        'search_reader': search_reader,
    })


@login_required
@role_required(['librarian', 'director'])
def export_library_excel(request):
    """
    Exports library fund and active borrowings into an Excel file (.xlsx).
    """
    wb = openpyxl.Workbook()

    # Sheet 1: Active Borrowings & Debtors
    ws1 = wb.active
    ws1.title = "Выданные книги"
    ws1.append(["№ Выдачи", "Читатель", "Роль", "Книга", "Инв. номер", "Дата выдачи", "Срок сдачи", "Просрочено (дней)", "Статус"])

    today = timezone.now().date()
    for b in BorrowRecord.objects.filter(return_date__isnull=True).select_related('reader', 'copy__book').order_by('due_date'):
        days = b.days_overdue
        status = f"ПРОСРОЧЕНО ({days} дн.)" if b.is_overdue else "На руках"
        ws1.append([
            b.id,
            b.reader.display_name,
            b.reader.role_title,
            b.copy.book.title,
            b.copy.inventory_number,
            b.borrow_date.strftime('%d.%m.%Y'),
            b.due_date.strftime('%d.%m.%Y'),
            days,
            status
        ])

    # Sheet 2: Books Catalog & Fund
    ws2 = wb.create_sheet(title="Книжный фонд")
    ws2.append(["ID", "Название", "Автор", "Жанр", "Класс", "Язык", "Год", "ISBN", "Всего экз.", "В наличии", "Выдано"])

    for book in Book.objects.all().select_related('genre').order_by('title'):
        ws2.append([
            book.id,
            book.title,
            book.author,
            book.genre.name if book.genre else "—",
            book.grade_level or "Все",
            book.get_language_display(),
            book.publication_year or "—",
            book.isbn,
            book.total_copies_count,
            book.available_copies_count,
            book.borrowed_copies_count
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="library_report_{today.isoformat()}.xlsx"'
    wb.save(response)
    return response


@login_required
@role_required(['librarian', 'director'])
def import_catalog_view(request):
    """
    Mass catalog import from CSV or Excel (.xlsx).
    Expected columns: title, author, genre, grade_level, language, publication_year, isbn, description, copies_count
    """
    if request.method == 'POST' and request.FILES.get('catalog_file'):
        file = request.FILES['catalog_file']
        filename = file.name.lower()
        imported = 0

        try:
            rows = []
            if filename.endswith('.csv'):
                decoded = file.read().decode('utf-8-sig', errors='ignore')
                reader = csv.reader(io.StringIO(decoded), delimiter=',' if ',' in decoded[:100] else ';')
                headers = [h.strip().lower() for h in next(reader)]
                for line in reader:
                    if line and any(line):
                        rows.append(dict(zip(headers, [v.strip() for v in line])))
            elif filename.endswith('.xlsx'):
                wb = openpyxl.load_workbook(file, data_only=True)
                sheet = wb.active
                iter_rows = sheet.iter_rows(values_only=True)
                headers = [str(h).strip().lower() if h else '' for h in next(iter_rows)]
                for r in iter_rows:
                    if any(r):
                        row_dict = {headers[i]: str(r[i]).strip() if r[i] is not None else '' for i in range(len(headers))}
                        rows.append(row_dict)
            else:
                messages.error(request, "Поддерживаются файлы CSV и Excel (.xlsx).")
                return redirect('library:librarian_cabinet')

            for r in rows:
                title = r.get('title') or r.get('название')
                author = r.get('author') or r.get('автор')
                if not title or not author:
                    continue

                genre_name = r.get('genre') or r.get('жанр')
                genre = None
                if genre_name:
                    genre, _ = Genre.objects.get_or_create(name=genre_name)

                grade_raw = r.get('grade_level') or r.get('класс') or ''
                grade_digits = ''.join([c for c in grade_raw if c.isdigit()])
                grade = int(grade_digits) if grade_digits else None

                lang = r.get('language') or r.get('язык') or 'ru'
                if 'кырг' in lang.lower() or 'kg' in lang.lower():
                    lang = 'kg'
                elif 'англ' in lang.lower() or 'en' in lang.lower():
                    lang = 'en'
                else:
                    lang = 'ru'

                year_raw = r.get('publication_year') or r.get('год') or ''
                year_digits = ''.join([c for c in year_raw if c.isdigit()])
                year = int(year_digits) if year_digits else None

                isbn = r.get('isbn') or ''
                desc = r.get('description') or r.get('описание') or ''
                copies_count_raw = r.get('copies_count') or r.get('экземпляры') or '1'
                copies_cnt = int(''.join([c for c in copies_count_raw if c.isdigit()]) or '1')

                book, _ = Book.objects.get_or_create(
                    title=title, author=author,
                    defaults={
                        'genre': genre,
                        'grade_level': grade,
                        'language': lang,
                        'publication_year': year,
                        'isbn': isbn,
                        'description': desc
                    }
                )

                # Add copies
                existing_copies = book.copies.count()
                for c_idx in range(1, copies_cnt + 1):
                    inv = f"INV-{book.id}-{existing_copies + c_idx:03d}"
                    BookCopy.objects.get_or_create(book=book, inventory_number=inv, defaults={'status': 'available'})

                imported += 1

            messages.success(request, f"Успешно импортировано книг в каталог: {imported}!")
            return redirect('library:librarian_cabinet')

        except Exception as e:
            messages.error(request, f"Ошибка при импорте каталога: {str(e)}")

    return render(request, 'library/import_catalog.html')


@login_required
@role_required(['teacher'])
def recommend_book_view(request, book_id):
    """
    Teacher can recommend a book to their classes.
    """
    book = get_object_or_404(Book, id=book_id)
    if request.method == 'POST':
        class_id = request.POST.get('target_class')
        note = request.POST.get('note', '').strip()
        target_class = SchoolClass.objects.filter(id=class_id).first() if class_id else None

        BookRecommendation.objects.create(
            teacher=request.user,
            book=book,
            target_class=target_class,
            note=note or "Рекомендовано учителем для углубленного чтения."
        )
        messages.success(request, f"Книга «{book.title}» успешно рекомендована!")
        return redirect('library:book_detail', book_id=book.id)

    classes = SchoolClass.objects.all().order_by('number', 'letter')
    return render(request, 'library/recommend_book.html', {'book': book, 'classes': classes})
